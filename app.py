#!/usr/bin/env python3
"""Local research collector for consented WeChat public-account and Moments data."""
from __future__ import annotations

import cgi
import hashlib
import html
import io
import json
import re
import sqlite3
import subprocess
import sys
import threading
import urllib.request
import webbrowser
import zipfile
from datetime import datetime, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from docx import Document
from lxml import html as lxml_html
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

SOURCE_ROOT = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
if getattr(sys, "frozen", False) and sys.platform == "darwin":
    ROOT = Path.home() / "Library" / "Application Support" / "微信研究资料整理器"
else:
    ROOT = Path(__file__).resolve().parent
ROOT.mkdir(parents=True, exist_ok=True)
DATA = ROOT / "data"
UPLOADS = DATA / "uploads"
DB = DATA / "research.db"
CONFIG = ROOT / "categories.json"
EXPORTS = ROOT / "exports"
MAX_UPLOAD = 20 * 1024 * 1024

if not CONFIG.exists():
    bundled_config = SOURCE_ROOT / "categories.json"
    if bundled_config.exists():
        CONFIG.write_bytes(bundled_config.read_bytes())


def now() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def db() -> sqlite3.Connection:
    DATA.mkdir(exist_ok=True)
    UPLOADS.mkdir(exist_ok=True)
    EXPORTS.mkdir(exist_ok=True)
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    con.executescript("""
    CREATE TABLE IF NOT EXISTS participants(
      code TEXT PRIMARY KEY, consent_ref TEXT NOT NULL, consent_date TEXT NOT NULL,
      notes TEXT DEFAULT '', active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS items(
      id INTEGER PRIMARY KEY AUTOINCREMENT, participant_code TEXT NOT NULL,
      source_type TEXT NOT NULL, source_ref TEXT NOT NULL, published_at TEXT DEFAULT '',
      title TEXT DEFAULT '', content TEXT NOT NULL, media_path TEXT DEFAULT '',
      auto_tags TEXT DEFAULT '[]', manual_tags TEXT DEFAULT '[]', memo TEXT DEFAULT '',
      review_status TEXT DEFAULT '待复核', content_hash TEXT UNIQUE NOT NULL,
      collected_at TEXT NOT NULL,
      FOREIGN KEY(participant_code) REFERENCES participants(code));
    """)
    return con


def clean_text(value: str) -> str:
    return re.sub(r"\n{3,}", "\n\n", re.sub(r"[ \t]+", " ", value or "")).strip()


def classify(text: str) -> list[str]:
    config = json.loads(CONFIG.read_text(encoding="utf-8"))
    low = text.lower()
    scores = []
    for category, words in config.items():
        hits = sum(low.count(word.lower()) for word in words)
        if hits:
            scores.append((hits, category))
    return [name for _, name in sorted(scores, reverse=True)[:6]] or ["未分类"]


def fetch_article(url: str) -> tuple[str, str, str]:
    if not re.match(r"^https://mp\.weixin\.qq\.com/s(?:/|\?)", url):
        raise ValueError("仅接受 mp.weixin.qq.com/s 开头的公众号文章链接")
    req = urllib.request.Request(url, headers={
        "User-Agent": "Mozilla/5.0 (research archiving; user-submitted URL)",
        "Accept-Language": "zh-CN,zh;q=0.9",
    })
    with urllib.request.urlopen(req, timeout=20) as res:
        if int(res.headers.get("Content-Length", "0") or 0) > 5 * 1024 * 1024:
            raise ValueError("页面过大")
        raw = res.read(5 * 1024 * 1024)
    tree = lxml_html.fromstring(raw)
    title = "".join(tree.xpath("//meta[@property='og:title']/@content") or tree.xpath("//title/text()"))
    nodes = tree.xpath("//*[@id='js_content']")
    if not nodes:
        raise ValueError("微信未返回正文；请将文章另存为 HTML/PDF，或复制正文后导入")
    content = clean_text(nodes[0].text_content())
    date = "".join(tree.xpath("//*[@id='publish_time']/text()") or tree.xpath("//meta[@property='article:published_time']/@content"))
    if len(content) < 30:
        raise ValueError("正文为空或遇到访问验证；请改用文件/文本导入")
    return clean_text(title), content, clean_text(date)


def extract_file(filename: str, payload: bytes) -> tuple[str, str]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md"}:
        return Path(filename).stem, payload.decode("utf-8", errors="replace")
    if suffix in {".html", ".htm"}:
        tree = lxml_html.fromstring(payload)
        for bad in tree.xpath("//script|//style|//noscript"):
            bad.drop_tree()
        title = "".join(tree.xpath("//title/text()")) or Path(filename).stem
        nodes = tree.xpath("//*[@id='js_content']")
        return clean_text(title), clean_text((nodes[0] if nodes else tree).text_content())
    if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        UPLOADS.mkdir(parents=True, exist_ok=True)
        digest = hashlib.sha256(payload).hexdigest()[:16]
        path = UPLOADS / f"{digest}{suffix}"
        path.write_bytes(payload)
        note = macos_vision_ocr(path) if sys.platform == "darwin" else tesseract_ocr(path)
        return Path(filename).stem, note
    raise ValueError("支持 TXT、MD、HTML、PNG、JPG、WEBP；PDF建议先复制文字或转为图片")


def macos_vision_ocr(path: Path) -> str:
    """Use Apple's built-in, offline Vision OCR in the packaged macOS app."""
    try:
        from Foundation import NSURL
        from Vision import VNImageRequestHandler, VNRecognizeTextRequest, VNRequestTextRecognitionLevelAccurate
        request = VNRecognizeTextRequest.alloc().init()
        request.setRecognitionLevel_(VNRequestTextRecognitionLevelAccurate)
        request.setRecognitionLanguages_(["zh-Hans", "en-US"])
        request.setUsesLanguageCorrection_(True)
        url = NSURL.fileURLWithPath_(str(path))
        handler = VNImageRequestHandler.alloc().initWithURL_options_(url, {})
        ok, error = handler.performRequests_error_([request], None)
        if not ok:
            raise RuntimeError(str(error))
        lines = []
        for observation in request.results() or []:
            candidates = observation.topCandidates_(1)
            if candidates:
                lines.append(str(candidates[0].string()))
        return clean_text("\n".join(lines)) or "[未识别出文字，请人工粘贴或录入]"
    except Exception as exc:
        return f"[macOS OCR 未成功：{exc}；请人工粘贴或录入]"


def tesseract_ocr(path: Path) -> str:
    try:
        langs = subprocess.run(["tesseract", "--list-langs"], capture_output=True, text=True).stdout
        lang = "chi_sim+eng" if "chi_sim" in langs else "eng"
        run = subprocess.run(["tesseract", str(path), "stdout", "-l", lang], capture_output=True, text=True)
        note = run.stdout.strip()
        if "chi_sim" not in langs:
            note = "[未安装中文OCR语言包；请人工校对或粘贴文字]\n" + note
        return note
    except FileNotFoundError:
        return "[未安装 OCR；请人工粘贴或录入]"


def add_item(code: str, source_type: str, source_ref: str, title: str, content: str,
             published_at: str = "", media_path: str = "") -> int:
    content = clean_text(content)
    if not content:
        raise ValueError("没有可保存的正文")
    con = db()
    p = con.execute("SELECT active FROM participants WHERE code=?", (code,)).fetchone()
    if not p or not p["active"]:
        raise ValueError("受访者不存在、未登记同意，或已撤回")
    digest = hashlib.sha256((code + source_ref + content).encode()).hexdigest()
    try:
        cur = con.execute("""INSERT INTO items(participant_code,source_type,source_ref,published_at,title,
          content,media_path,auto_tags,content_hash,collected_at) VALUES(?,?,?,?,?,?,?,?,?,?)""",
          (code, source_type, source_ref, published_at, title, content, media_path,
           json.dumps(classify(title + "\n" + content), ensure_ascii=False), digest, now()))
        con.commit()
        return cur.lastrowid
    except sqlite3.IntegrityError as e:
        raise ValueError("这条资料已经导入过") from e


def esc(x) -> str:
    return html.escape(str(x or ""), quote=True)


STYLE = """
body{font:15px/1.55 system-ui;margin:0;background:#f5f7f8;color:#17212b}.wrap{max-width:1080px;margin:auto;padding:24px}
nav{background:#087f5b;color:white;padding:15px}nav a{color:white;margin-right:18px}section,.card{background:white;padding:18px;margin:16px 0;border-radius:10px;box-shadow:0 1px 5px #0001}
input,select,textarea,button{font:inherit;padding:8px;margin:4px 2px;border:1px solid #bbb;border-radius:6px}textarea{width:96%}button{background:#087f5b;color:white;border:0;cursor:pointer}
table{width:100%;border-collapse:collapse}th,td{text-align:left;vertical-align:top;border-bottom:1px solid #ddd;padding:8px}.muted{color:#68737d}.warn{background:#fff3bf;padding:10px;border-radius:6px}
"""


def page(body: str, msg: str = "") -> bytes:
    banner = f'<p class="warn">{esc(msg)}</p>' if msg else ""
    return f'''<!doctype html><meta charset="utf-8"><title>微信研究资料整理器</title><style>{STYLE}</style>
    <nav><div class="wrap"><b>微信研究资料整理器</b>　<a href="/">采集</a><a href="/items">检索与编码</a><a href="/participants">知情同意</a><a href="/export">导出</a></div></nav>
    <main class="wrap">{banner}{body}</main>'''.encode()


class Handler(BaseHTTPRequestHandler):
    def send_page(self, body, msg="", status=200):
        data = page(body, msg); self.send_response(status); self.send_header("Content-Type", "text/html; charset=utf-8"); self.end_headers(); self.wfile.write(data)

    def fields(self):
        ctype, pdict = cgi.parse_header(self.headers.get("Content-Type", ""))
        if ctype == "multipart/form-data":
            return cgi.FieldStorage(fp=self.rfile, headers=self.headers, environ={"REQUEST_METHOD":"POST", "CONTENT_TYPE":self.headers["Content-Type"]})
        length = min(int(self.headers.get("Content-Length", 0)), MAX_UPLOAD)
        from urllib.parse import parse_qs
        return {k: v[-1] for k,v in parse_qs(self.rfile.read(length).decode()).items()}

    def getv(self, f, key, default=""):
        if isinstance(f, dict): return f.get(key, default)
        return f.getfirst(key, default)

    def do_GET(self):
        from urllib.parse import urlparse, parse_qs
        p = urlparse(self.path); q = parse_qs(p.query)
        if p.path == "/":
            codes = [r[0] for r in db().execute("SELECT code FROM participants WHERE active=1 ORDER BY code")]
            opts = ''.join(f'<option>{esc(c)}</option>' for c in codes)
            body = f'''<h1>采集资料</h1><p class="muted">先登记知情同意。公众号按链接采集；朋友圈请上传截图/HTML，或粘贴文字。</p>
            <section><h2>公众号文章链接</h2><form method="post" action="/add-url"><select name="code">{opts}</select><input size="65" name="url" placeholder="https://mp.weixin.qq.com/s/…" required><button>采集</button></form></section>
            <section><h2>朋友圈或其他材料</h2><form method="post" action="/add-text"><select name="code">{opts}</select><input name="title" placeholder="标题/日期简述"><input name="source_ref" placeholder="来源说明，如朋友圈截图 2026-09-20"><textarea rows="10" name="content" placeholder="粘贴文字；截图可在下一栏上传"></textarea><button>保存文字</button></form>
            <form method="post" action="/add-file" enctype="multipart/form-data"><select name="code">{opts}</select><input type="file" name="file" required><button>上传并提取</button></form></section>'''
            return self.send_page(body)
        if p.path == "/participants":
            rows = db().execute("SELECT * FROM participants ORDER BY code").fetchall()
            table=''.join(f'<tr><td>{esc(r["code"])}</td><td>{esc(r["consent_ref"])}</td><td>{esc(r["consent_date"])}</td><td>{"有效" if r["active"] else "已撤回"}</td><td><form method="post" action="/withdraw"><input type="hidden" name="code" value="{esc(r["code"])}"><button>撤回并删除资料</button></form></td></tr>' for r in rows)
            body=f'''<h1>知情同意登记</h1><section><form method="post" action="/participant"><input name="code" placeholder="受访者代号，如 W01" required><input name="consent_ref" placeholder="同意书编号/存放位置" required><input type="date" name="consent_date" required><input name="notes" placeholder="范围或限制"><button>登记</button></form></section><section><table><tr><th>代号</th><th>同意依据</th><th>日期</th><th>状态</th><th>操作</th></tr>{table}</table></section>'''
            return self.send_page(body)
        if p.path == "/items":
            term=q.get("q",[""])[0]; tag=q.get("tag",[""])[0]
            sql="SELECT * FROM items WHERE (title LIKE ? OR content LIKE ? OR participant_code LIKE ?) AND (auto_tags LIKE ? OR manual_tags LIKE ?) ORDER BY id DESC"
            like=f"%{term}%"; t=f"%{tag}%"; rows=db().execute(sql,(like,like,like,t,t)).fetchall()
            cards=''.join(f'''<div class="card"><b>#{r["id"]} {esc(r["participant_code"])}｜{esc(r["source_type"])}｜{esc(r["title"])}</b><p class="muted">{esc(r["published_at"] or r["collected_at"])} · {esc(r["source_ref"])}</p><p>{esc(r["content"][:800])}</p><form method="post" action="/code"><input type="hidden" name="id" value="{r["id"]}"><input size="45" name="manual_tags" value="{esc(', '.join(json.loads(r['manual_tags']))) }" placeholder="人工标签，逗号分隔"><select name="review_status"><option>{esc(r["review_status"])}</option><option>已复核</option><option>待复核</option><option>排除</option></select><input size="45" name="memo" value="{esc(r["memo"])}" placeholder="研究备忘"><button>更新编码</button></form><small>自动标签：{esc(', '.join(json.loads(r['auto_tags'])))}</small></div>''' for r in rows)
            return self.send_page(f'''<h1>检索与编码</h1><form><input name="q" value="{esc(term)}" placeholder="全文/受访者检索"><input name="tag" value="{esc(tag)}" placeholder="标签"><button>检索</button></form>{cards or '<p>暂无资料。</p>'}''')
        if p.path == "/export":
            return self.send_page('''<h1>导出</h1><p>导出前请完成必要的人工复核。文件仅含受访者代号，不包含真实姓名。</p><form method="post" action="/make-export"><button>生成 Excel + Word + 数据库备份</button></form>''')
        if p.path.startswith("/download/"):
            name=Path(p.path).name; target=EXPORTS/name
            if not target.exists(): return self.send_error(404)
            self.send_response(200); self.send_header("Content-Type","application/zip"); self.send_header("Content-Disposition",f'attachment; filename="{name}"'); self.end_headers(); self.wfile.write(target.read_bytes()); return
        self.send_error(404)

    def do_POST(self):
        try:
            f=self.fields(); path=self.path
            if path=="/participant":
                con=db(); con.execute("INSERT OR REPLACE INTO participants(code,consent_ref,consent_date,notes,active,created_at) VALUES(?,?,?,?,1,?)",(self.getv(f,"code").strip(),self.getv(f,"consent_ref"),self.getv(f,"consent_date"),self.getv(f,"notes"),now())); con.commit(); return self.redirect("/participants","已登记")
            if path=="/withdraw":
                code=self.getv(f,"code"); con=db(); paths=[r[0] for r in con.execute("SELECT media_path FROM items WHERE participant_code=?",(code,)) if r[0]]; con.execute("DELETE FROM items WHERE participant_code=?",(code,)); con.execute("UPDATE participants SET active=0 WHERE code=?",(code,)); con.commit()
                for x in paths:
                    try: Path(x).unlink(missing_ok=True)
                    except OSError: pass
                return self.redirect("/participants","已撤回并删除该代号的资料")
            if path=="/add-url":
                url=self.getv(f,"url").strip(); title,content,date=fetch_article(url); add_item(self.getv(f,"code"),"公众号",url,title,content,date); return self.redirect("/items","文章已采集，请人工复核")
            if path=="/add-text":
                add_item(self.getv(f,"code"),"朋友圈/文本",self.getv(f,"source_ref") or "人工粘贴",self.getv(f,"title"),self.getv(f,"content")); return self.redirect("/items","资料已保存")
            if path=="/add-file":
                item=f["file"]; payload=item.file.read(MAX_UPLOAD+1)
                if len(payload)>MAX_UPLOAD: raise ValueError("文件超过 20MB")
                title,content=extract_file(Path(item.filename).name,payload); digest=hashlib.sha256(payload).hexdigest()[:16]; suffix=Path(item.filename).suffix.lower(); media=str(UPLOADS/f"{digest}{suffix}") if suffix in {'.png','.jpg','.jpeg','.webp'} else ""
                add_item(self.getv(f,"code"),"朋友圈/文件",Path(item.filename).name,title,content,media_path=media); return self.redirect("/items","文件已提取，请核对 OCR/正文")
            if path=="/code":
                tags=[x.strip() for x in re.split('[,，]',self.getv(f,"manual_tags")) if x.strip()]
                con=db(); con.execute("UPDATE items SET manual_tags=?,review_status=?,memo=? WHERE id=?",(json.dumps(tags,ensure_ascii=False),self.getv(f,"review_status"),self.getv(f,"memo"),int(self.getv(f,"id")))); con.commit(); return self.redirect("/items","编码已更新")
            if path=="/make-export":
                name=make_export(); return self.redirect("/export",f'已生成：<a href="/download/{name}">{name}</a>',raw=True)
            self.send_error(404)
        except Exception as e:
            self.send_page("<h1>操作未完成</h1><p>请返回修改输入。</p>",str(e),400)

    def redirect(self, path, msg, raw=False):
        from urllib.parse import quote
        # Keep messages plain; export link is rendered on a small response page.
        if raw:
            self.send_page(f'<p>{msg}</p><p><a href="{path}">返回</a></p>'); return
        self.send_response(303); self.send_header("Location",path+"?msg="+quote(msg)); self.end_headers()

    def log_message(self, fmt, *args):
        print("[%s] %s" % (self.log_date_time_string(), fmt % args))


def make_export() -> str:
    rows=db().execute("SELECT * FROM items ORDER BY participant_code,id").fetchall()
    stamp=datetime.now().strftime("%Y%m%d_%H%M%S")
    xlsx=EXPORTS/f"研究编码表_{stamp}.xlsx"; docx=EXPORTS/f"个案档案_{stamp}.docx"; zpath=EXPORTS/f"微信研究资料导出_{stamp}.zip"
    wb=Workbook(); ws=wb.active; ws.title="编码资料"; headers=["ID","受访者代号","来源类型","来源","发布日期","标题","原文","自动标签","人工标签","研究备忘","复核状态","采集时间"]
    ws.append(headers)
    for c in ws[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="087F5B")
    for r in rows: ws.append([r["id"],r["participant_code"],r["source_type"],r["source_ref"],r["published_at"],r["title"],r["content"],", ".join(json.loads(r["auto_tags"])),", ".join(json.loads(r["manual_tags"])),r["memo"],r["review_status"],r["collected_at"]])
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    for col,w in {"A":8,"B":14,"C":16,"D":40,"E":18,"F":30,"G":80,"H":25,"I":25,"J":40,"K":12,"L":22}.items(): ws.column_dimensions[col].width=w
    wb.save(xlsx)
    doc=Document(); doc.add_heading("微信研究资料个案档案",0); doc.add_paragraph(f"生成时间：{now()}。资料仅以受访者代号组织，引用前须复核原文与知情同意范围。")
    for code in sorted({r['participant_code'] for r in rows}):
        doc.add_heading(code,level=1)
        for r in [x for x in rows if x['participant_code']==code]:
            doc.add_heading(r['title'] or f"资料 #{r['id']}",level=2); doc.add_paragraph(f"来源：{r['source_type']}｜{r['source_ref']}｜{r['published_at'] or r['collected_at']}"); doc.add_paragraph(r['content']); doc.add_paragraph("标签："+", ".join(json.loads(r['auto_tags'])+json.loads(r['manual_tags']))); doc.add_paragraph("备忘："+r['memo'])
    doc.save(docx)
    with zipfile.ZipFile(zpath,"w",zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx,xlsx.name); z.write(docx,docx.name); z.write(DB,"research.db"); z.write(CONFIG,"categories.json")
    return zpath.name


if __name__ == "__main__":
    db().close(); port=int(sys.argv[1]) if len(sys.argv)>1 else 8765
    print(f"Open http://127.0.0.1:{port}  (Ctrl+C to stop)")
    if getattr(sys, "frozen", False):
        threading.Timer(1.0, lambda: webbrowser.open(f"http://127.0.0.1:{port}")).start()
    ThreadingHTTPServer(("127.0.0.1",port),Handler).serve_forever()
